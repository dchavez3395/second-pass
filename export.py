"""
Second Pass — fill the agency's Project Accessibility Sheets template for one site.

    python export.py exports/<site>.input.json     -> exports/<site>.xlsx (path printed last)

The review server writes the input file: the site URL, a status and notes per
criterion (set by the reviewer, or implied by the review), and the task list
(logged issues plus one row per issue with confirmed elements). This script only
copies the template and writes those into "Review WCAG V2.2" and "Task List".

The template path can be overridden with SECOND_PASS_TEMPLATE.
"""
import json
import os
import re
import shutil
import sys
from datetime import date

import openpyxl

ROOT = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.environ.get(
    "SECOND_PASS_TEMPLATE",
    r"D:\Downloads\Copy of [TEMPLATE] Project Accessibility Sheets (1).xlsx",
)


def main():
    if len(sys.argv) < 2:
        sys.exit("usage: export.py <input.json>")
    with open(sys.argv[1], encoding="utf-8") as fh:
        data = json.load(fh)
    site = data["site"]
    if not os.path.exists(TEMPLATE):
        sys.exit(f"template not found: {TEMPLATE}")

    out = os.path.join(ROOT, "exports", f"{site}.xlsx")
    shutil.copyfile(TEMPLATE, out)
    wb = openpyxl.load_workbook(out)

    # ---- Review WCAG V2.2: A2 = URL; per criterion row, D = status, E = notes
    ws = wb["Review WCAG V2.2"]
    ws["A2"] = f"{data.get('url', '')}  (Second Pass review, {date.today().isoformat()})"
    filled = 0
    for row in ws.iter_rows(min_row=4):
        level, crit = row[0].value, row[1].value
        if level not in ("A", "AA", "AAA") or not crit:
            continue
        m = re.match(r"(\d\.\d\.\d+)", str(crit))
        if not m or m.group(1) not in data["criteria"]:
            continue
        rec = data["criteria"][m.group(1)]
        row[3].value = rec.get("status") or "Not Evaluated"
        if rec.get("notes"):
            row[4].value = rec["notes"]
        filled += 1

    # ---- Task List: one row per issue, after the header
    tl = wb["Task List"]
    r = 2
    while tl.cell(row=r, column=3).value:
        r += 1
    names = data.get("criterionNames", {})
    tasks = sorted(data["tasks"], key=lambda t: -int(t.get("severity") or 0))
    for t in tasks:
        tl.cell(row=r, column=1).value = int(t.get("severity") or 2)
        tl.cell(row=r, column=2).value = "@ _ Name"
        tl.cell(row=r, column=3).value = t.get("shortname", "")
        tl.cell(row=r, column=4).value = names.get(t.get("criterion", ""), t.get("criterion", ""))
        tl.cell(row=r, column=5).value = t.get("status", "Not Started")
        tl.cell(row=r, column=6).value = t.get("notes", "")
        tl.cell(row=r, column=7).value = t.get("reference", "")
        r += 1

    wb.save(out)
    print(f"{filled} criteria, {len(tasks)} issues -> {out}")
    print(out)


if __name__ == "__main__":
    main()
